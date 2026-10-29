"""Career Skill Gap Analyzer: interactive dashboard.
Run:  streamlit run app.py        (after:  python build_model.py ...)
"""
import io, os, re, sys, joblib
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
for p in (HERE, os.path.dirname(HERE)):
    sys.path.insert(0, p)
import auth
import skillgap_core as K
import career_skill_gap_analysis as A
from skillgap_core import SKILLS, SIDX, ROLE_NAMES, CATEGORY

st.set_page_config(page_title="Career Skill Gap Analyzer", page_icon="🎯", layout="wide")
BLUE, RED, GREY = "#1f4e79", "#c0504d", "#8a8a8a"


@st.cache_resource
def load_bundle():
    path = os.path.join(HERE, "model.pkl")
    if not os.path.exists(path):
        st.error("model.pkl not found. Run:  python build_model.py --postings data/postings.csv --cache outputs/cache")
        st.stop()
    return joblib.load(path)


USER = auth.gate()
B = load_bundle()
IMP, MODEL, ALPHA, POP = B["imp_all"], B["M"], B["alpha"], B["pop"]
SAL = {r["skill"]: r for r in B["salary"]}


# ------------------------------------------------------------------ resume reading
def read_resume(upload):
    name = upload.name.lower()
    data = upload.read()
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


LEARN_CUE = re.compile(r"((currently|now|still|am|also)\s+(learning|studying|exploring)|\bexploring\b|\bbeginner\b|want(ing)? to learn|will learn|basic knowledge of|learning to)", re.I)


def split_learning(text):
    """Separate sentences that only say the person is learning a skill from the rest."""
    keep, learn = [], []
    for seg in re.split(r"(?<=[.!?])\s+|\n", text):
        (learn if LEARN_CUE.search(seg) and len(seg) < 200 else keep).append(seg)
    return "\n".join(keep), "\n".join(learn)


def detect_skills(text):
    main, learn = split_learning(text)
    def run(t):
        found = set(K.extract_text(t, "alias"))
        for item in re.split(r"[,\n;|\u2022\u00b7/()]+", t):
            item = item.strip(" -:\t")
            if 2 <= len(item) <= 40:
                m = K.map_item(item, "fuzzy")
                if m:
                    found.add(m)
        return found
    held = run(main)
    learning = sorted(run(learn) - held)
    return sorted(held), learning


# ------------------------------------------------------------------ analysis
def analyse(skills, role_i):
    x = K.vec(skills)[None, :]
    rd_all = A.readiness(x, IMP)[0]
    proba = MODEL.proba(x, ALPHA)[0]
    req = [j for j in np.where(IMP[role_i] > 0)[0]]
    have = [SKILLS[j] for j in req if x[0, j] > 0]
    gap = sorted([SKILLS[j] for j in req if x[0, j] == 0], key=lambda s: -IMP[role_i, SIDX[s]])
    seq = A.plan(skills, role_i, IMP, "greedy", np.random.default_rng(0), POP)
    return x[0], rd_all, proba, have, gap, seq


def radar(x, role_i):
    cats = list(dict.fromkeys(CATEGORY.values()))
    vals = []
    for c in cats:
        rc = [s for s in SKILLS if CATEGORY[s] == c and IMP[role_i, SIDX[s]] > 0]
        tot = sum(IMP[role_i, SIDX[s]] for s in rc)
        vals.append(sum(IMP[role_i, SIDX[s]] for s in rc if x[SIDX[s]] > 0) / tot if tot > 0 else np.nan)
    ang = np.linspace(0, 2 * np.pi, len(cats), endpoint=False)
    v = np.nan_to_num(vals, nan=0.0)
    fig = plt.figure(figsize=(4.2, 4.2)); ax = fig.add_subplot(111, polar=True)
    ax.plot(np.r_[ang, ang[0]], np.r_[v, v[0]], color=BLUE, lw=2); ax.fill(np.r_[ang, ang[0]], np.r_[v, v[0]], color=BLUE, alpha=0.25)
    ax.plot(np.r_[ang, ang[0]], np.ones(len(cats) + 1), color=RED, lw=1, ls="--")
    ax.set_xticks(ang); ax.set_xticklabels([c.replace(" & ", "\n& ") for c in cats], fontsize=8)
    ax.set_yticks([0.5, 1.0]); ax.set_yticklabels(["50%", "100%"], fontsize=7); ax.set_ylim(0, 1.05)
    ax.set_title("Coverage of the role's requirement, by skill group\n(dashed = fully covered; roles with no skills in a group show 0)", fontsize=8)
    return fig



def show_history():
    h = auth.get_history(USER)
    if not h:
        st.info("No saved analyses yet. Analyse a resume and press 'Save this analysis to my history'.")
        return
    hd = pd.DataFrame(h)
    st.dataframe(hd, hide_index=True)
    if len(hd) > 1:
        fig, ax = plt.subplots(figsize=(7, 2.4))
        ax.plot(range(1, len(hd) + 1), hd["readiness"], marker="o", color=BLUE)
        ax.set_xlabel("Saved analysis #"); ax.set_ylabel("Readiness"); ax.set_ylim(0, 100); ax.grid(lw=0.3)
        st.pyplot(fig)


# ------------------------------------------------------------------ sidebar
st.sidebar.title("🎯 Skill Gap Analyzer")
st.sidebar.caption(f"Role requirements learned from {B['n_postings']:,} real LinkedIn IT postings, {K.R_N} roles, {K.S_N} skills.")
up = st.sidebar.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
pasted = st.sidebar.text_area("…or paste resume text / a skill list", height=140)
text = ""
if up is not None:
    try:
        text = read_resume(up)
    except Exception as e:
        st.sidebar.error(f"Could not read the file: {e}")
elif pasted.strip():
    text = pasted

st.title("Career Skill Gap Analyzer and Job Recommendation")
if not text.strip():
    st.info("Upload a resume or paste text in the sidebar to start. The app finds your skills, scores your readiness for IT roles, "
            "shows the skill gap and builds a learning roadmap ordered by gain per study hour.")
    st.subheader("My history")
    show_history()
    with st.expander("How it works"):
        st.markdown("1. NLP extraction maps the text to a 133-skill taxonomy (aliases, typos, context).\n"
                    "2. Role importance weights come from real job postings (frequency × role specificity).\n"
                    "3. Readiness = importance-weighted share of a role's required skills you hold.\n"
                    "4. Role ranking = Random Forest + SVM probabilities blended with readiness.\n"
                    "5. The roadmap picks the missing skill with the best readiness gain per hour, adding prerequisites first.")
    st.stop()

detected, learning = detect_skills(text)
st.sidebar.success(f"{len(detected)} skills detected")
if learning:
    st.sidebar.info("Not counted (you say you are still learning): " + ", ".join(learning))
skills = st.sidebar.multiselect("Detected skills (add or remove to correct them)", SKILLS, default=detected)
if not skills:
    st.warning("No skills recognised. Add some in the sidebar list.")
    st.stop()

x0 = K.vec(skills)[None, :]
proba0 = MODEL.proba(x0, ALPHA)[0]
top3 = np.argsort(-proba0)[:3]
role_choice = st.sidebar.selectbox("Target role", ["Best match: " + ROLE_NAMES[top3[0]]] + ROLE_NAMES)
role_i = top3[0] if role_choice.startswith("Best match") else ROLE_NAMES.index(role_choice)
x, rd_all, proba, have, gap, seq = analyse(skills, role_i)
rd = rd_all[role_i] * 100

tab1, tab2, tab3, tab4, tab_h, tab5 = st.tabs(["Overview", "Skill gap", "Learning roadmap", "Salary insights", "My history", "About & limits"])

with tab1:
    c1, c2, c3 = st.columns(3)
    c1.metric("Target role", ROLE_NAMES[role_i])
    c2.metric("Readiness", f"{rd:.0f} / 100")
    c3.metric("Required skills held", f"{len(have)} of {len(have) + len(gap)}")
    st.progress(min(int(rd), 100))
    l, r = st.columns([1, 1])
    with l:
        st.subheader("Top 3 role matches")
        fig, ax = plt.subplots(figsize=(4.4, 2.2))
        ax.barh(range(3)[::-1], proba[top3], color=[BLUE, GREY, GREY])
        ax.set_yticks(range(3)[::-1]); ax.set_yticklabels([ROLE_NAMES[j] for j in top3], fontsize=9)
        ax.set_xlabel("Hybrid score"); ax.grid(axis="x", lw=0.3)
        st.pyplot(fig)
        st.caption("A top-3 list is more reliable than a single answer; choose the target role yourself in the sidebar.")
        rdf = pd.DataFrame({"Role": ROLE_NAMES, "Readiness": (rd_all * 100).round(0), "Match score": proba.round(3)}).sort_values("Match score", ascending=False)
        st.dataframe(rdf, hide_index=True, height=250)
    with r:
        st.subheader("Skill radar")
        st.pyplot(radar(x, role_i))

with tab2:
    st.subheader(f"Skill gap for {ROLE_NAMES[role_i]}")
    g1, g2 = st.columns(2)
    with g1:
        st.markdown("**Missing required skills** (most important first)")
        gdf = pd.DataFrame([{"Skill": s, "Group": CATEGORY[s], "Importance": round(float(IMP[role_i, SIDX[s]]), 2),
                             "Study hours": A.hours_of(s)} for s in gap])
        if len(gdf):
            st.dataframe(gdf, hide_index=True, height=420)
        else:
            st.success("No required skills missing.")
    with g2:
        st.markdown("**Required skills you already hold**")
        hdf = pd.DataFrame([{"Skill": s, "Group": CATEGORY[s], "Importance": round(float(IMP[role_i, SIDX[s]]), 2)}
                            for s in sorted(have, key=lambda s: -IMP[role_i, SIDX[s]])])
        if len(hdf):
            st.dataframe(hdf, hide_index=True, height=420)
        else:
            st.warning("None yet.")
    extra = [s for s in skills if IMP[role_i, SIDX[s]] == 0]
    if extra:
        st.caption("Other skills you have that this role does not list as required: " + ", ".join(extra))

with tab3:
    st.subheader("Learning roadmap (best readiness gain per study hour)")
    budget = st.slider("Study budget (hours)", 20, 600, 200, 20)
    steps, cum, cur = [], 0, rd
    for s in seq:
        h = A.hours_of(s)
        gain = float(IMP[role_i, SIDX[s]] / IMP[role_i].sum() * 100)
        cum += h; cur += gain
        steps.append(dict(Step=len(steps) + 1, Skill=s, Group=CATEGORY[s], Hours=h, **{"Cumulative hours": cum}, **{"Readiness after": round(min(cur, 100), 1)},
                          **{"Within budget": "yes" if cum <= budget else ""}))
    sdf = pd.DataFrame(steps)
    within = sdf[sdf["Cumulative hours"] <= budget]
    after = float(within["Readiness after"].iloc[-1]) if len(within) else rd
    m1, m2, m3 = st.columns(3)
    m1.metric("Readiness now", f"{rd:.0f}")
    m2.metric(f"Readiness after {budget} h", f"{after:.0f}", delta=f"{after - rd:+.0f}")
    m3.metric("Skills in budget", len(within))
    if len(sdf):
        fig, ax = plt.subplots(figsize=(7, 2.4))
        ax.plot([0] + list(sdf["Cumulative hours"]), [rd] + list(sdf["Readiness after"]), color=BLUE, lw=2)
        ax.axvline(budget, color=RED, ls="--", lw=1); ax.set_xlabel("Cumulative study hours"); ax.set_ylabel("Readiness"); ax.grid(lw=0.3)
        st.pyplot(fig)
        st.dataframe(sdf, hide_index=True)
    else:
        st.success("Nothing to learn for this role.")
    st.caption("Learning hours and prerequisites are assumptions made by the project team, not measured values.")

with tab4:
    st.subheader("Pay associated with each skill (R / log-linear model)")
    st.caption("From salaried LinkedIn postings, controlling for role and experience level. This shows association, not a guaranteed raise.")
    if not SAL:
        st.info("Salary coefficients not found. Re-run build_model.py with --salary_csv.")
    else:
        rows = []
        for s in gap[:15]:
            if s in SAL:
                r_ = SAL[s]
                rows.append({"Missing skill": s, "Salary difference (%)": round(r_["pct"], 1), "t-value": round(r_["t"], 2),
                             "Significant (5%)": "yes" if abs(r_["t"]) > 1.96 else "no", "Postings": int(r_["n"])})
        if rows:
            st.markdown("**Missing skills from your gap list that have a salary estimate**")
            st.dataframe(pd.DataFrame(rows), hide_index=True)
        else:
            st.info("None of your missing skills is among the 30 skills with a salary estimate.")
        sd = pd.DataFrame(B["salary"]).sort_values("pct")
        fig, ax = plt.subplots(figsize=(7, 4.2))
        ax.barh(sd["skill"], sd["pct"], color=[RED if v < 0 else BLUE for v in sd["pct"]])
        ax.axvline(0, color="k", lw=0.6); ax.set_xlabel("Salary difference when the skill is mentioned (%)"); ax.tick_params(axis="y", labelsize=7); ax.grid(axis="x", lw=0.3)
        st.pyplot(fig)

with tab_h:
    st.subheader("My history")
    if st.button("Save this analysis to my history"):
        import datetime
        auth.add_history(USER, dict(date=datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), role=ROLE_NAMES[role_i],
                                    readiness=round(float(rd), 1), skills=len(skills), missing=len(gap)))
        st.success("Saved.")
    show_history()

with tab5:
    st.markdown(f"""
**What this app does:** extracts skills from your resume, compares them with requirements learned from {B['n_postings']:,} LinkedIn IT postings (about two weeks of data from 2024), and plans what to learn.

**Limits you should know**
- Only the {K.S_N} skills in the taxonomy are recognised. Correct the list in the sidebar if something is missing or wrong.
- Role recommendation is reliable as a top-3 list or at the family level (data / software / infrastructure), not as a single answer. On real resumes our top-1 accuracy was about 36%, top-3 about 72%.
- Readiness correlates only moderately with the matching score in the resume dataset (Spearman 0.29).
- Study hours and prerequisites are assumptions. Salary differences are associations from postings that list a salary.
- Your resume is processed in memory and is not stored. Only the summary you choose to save (date, role, score, counts) is kept. The login is a prototype: salted password hashes in a local file, no e-mail verification or password reset.
""")

# ------------------------------------------------------------------ report download
rep = [f"# Skill gap report: {ROLE_NAMES[role_i]}", "", f"Readiness: {rd:.0f}/100", f"Top matches: " + ", ".join(ROLE_NAMES[j] for j in top3), "",
       "## Skills detected", ", ".join(skills), "", "## Missing required skills (most important first)"]
rep += [f"- {s} ({A.hours_of(s)} h)" for s in gap] or ["- none"]
rep += ["", "## Roadmap"] + [f"{i + 1}. {s} ({A.hours_of(s)} h)" for i, s in enumerate(seq)]
st.sidebar.download_button("Download report (Markdown)", "\n".join(rep), file_name="skill_gap_report.md")
